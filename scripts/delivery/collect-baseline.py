"""Read-only source/reference evidence for BASE-01; never emits source rows."""
import collections
import datetime
import hashlib
import json
from pathlib import Path
import re
import subprocess
import sys

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "docs/delivery/baseline"


def git(*args):
    return subprocess.check_output(["git", *args], cwd=ROOT).decode("utf-8").strip()


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def write(name, value):
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / name).write_text(json.dumps(value, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def source_snapshot():
    # Identity only; do not read configured environments or databases.
    names = git("ls-files", "--cached", "--others", "--exclude-standard", "-z").split("\0")
    hashes = {}
    for name in names:
        path = ROOT / name
        if path.is_file() and (name.startswith(("src/", "prisma/", "docs/sheets/", "docs/pwa screens/")) or name in {"package.json", "package-lock.json", "next.config.ts", "vitest.config.ts"}):
            if path.suffix.lower() in {".db", ".sqlite", ".sqlite3"} or path.name.startswith(".env"):
                continue
            hashes[name] = sha(path)
    return {"head": git("rev-parse", "HEAD"), "branch": git("branch", "--show-current"), "files": hashes}


def branches():
    refs = {ref: git("rev-parse", ref) for ref in ["HEAD", "main", "origin/main", "origin/v2"]}
    main_only = sorted(set(git("ls-tree", "-r", "--name-only", "origin/main").splitlines()) - set(git("ls-tree", "-r", "--name-only", "HEAD").splitlines()))
    main = git("show", "origin/main:prisma/postgres/schema.prisma")
    current = (ROOT / "prisma/postgres/schema.prisma").read_text(encoding="utf-8")
    return {"refs": refs, "mergeBase": git("merge-base", "origin/main", "HEAD"), "divergence": git("rev-list", "--left-right", "--count", "origin/main...HEAD"), "mainOnlyPaths": main_only, "changedSharedPaths": git("diff", "--name-only", "--diff-filter=M", "origin/main", "HEAD").splitlines(), "mainOnlyModels": sorted(set(re.findall(r"^model (\w+)", main, re.M)) - set(re.findall(r"^model (\w+)", current, re.M))), "mainUniqueCommits": git("rev-list", "--cherry-pick", "--left-only", "origin/main...HEAD").splitlines()}


# Only fixed public field vocabulary leaves the extractor. No names, sheet titles,
# contacts, notes, URLs or formula expressions are exported.
VOCAB = {
    "name": "name", "student name": "student_name", "student's name": "student_name",
    "father name": "father_name", "father's name": "father_name", "father": "father_name",
    "guardian": "guardian", "guardian name": "guardian_name", "guardian contact": "guardian_contact",
    "phone": "phone", "phone number": "phone", "contact": "contact", "contact no": "contact",
    "contact number": "contact", "mobile": "mobile", "whatsapp": "whatsapp", "email": "email",
    "address": "address", "school": "school", "school name": "school", "class": "class",
    "age": "age", "dob": "date_of_birth", "date of birth": "date_of_birth", "gender": "gender",
    "cnic": "cnic", "city": "city", "park": "park", "park name": "park", "batch": "batch",
    "group": "group", "murabbi": "murabbi", "murabbi name": "murabbi_name", "muawin": "muawin",
    "date": "date", "day": "day", "week": "week", "topic": "topic", "title": "title",
    "content": "content", "category": "category", "lesson": "lesson", "session": "session",
    "activity": "activity", "status": "status", "attendance": "attendance", "present": "present",
    "absent": "absent", "total": "total", "remarks": "remarks", "notes": "notes",
    "caller": "caller", "response": "response", "call status": "call_status", "timestamp": "timestamp",
    "qualification": "qualification", "occupation": "occupation", "s no": "serial", "sr no": "serial",
    "registration id": "registration_id", "registration no": "registration_id", "id": "id",
    "disclaimer": "disclaimer", "link": "link", "resource": "resource", "name of student": "student_name",
}


def worksheets():
    import openpyxl
    books = []
    for file in sorted((ROOT / "docs/sheets").iterdir()):
        if file.suffix.lower() not in {".xls", ".xlsx"}:
            continue
        with file.open("rb") as stream:
            wb = openpyxl.load_workbook(stream, read_only=True, data_only=False)
            book = {"file": file.name, "sha256": sha(file), "sheets": []}
            for index, ws in enumerate(wb.worksheets, 1):
                # Some exports falsely declare A1:A1 while storing thousands of cells.
                ws.reset_dimensions()
                headers, dates, counts, used_rows, row_shapes = [], [], collections.Counter(), set(), collections.Counter()
                cols = collections.defaultdict(collections.Counter)
                for row in ws.iter_rows():
                    populated = []
                    for cell in row:
                        value = cell.value
                        if value is None:
                            continue
                        used_rows.add(cell.row)
                        populated.append(cell.column)
                        kind = "formula" if cell.data_type == "f" else "date" if isinstance(value, (datetime.datetime, datetime.date)) else "number" if isinstance(value, (int, float)) else "text"
                        counts[kind] += 1
                        cols[cell.column_letter][kind] += 1
                        if kind == "date" and cell.row <= 12:
                            dates.append(cell.coordinate)
                        if isinstance(value, str) and kind != "formula":
                            norm = re.sub(r"[.\s_:]+", " ", value.strip().lower()).strip()
                            if cell.row <= 12 and norm in VOCAB:
                                headers.append({"cell": cell.coordinate, "field": VOCAB[norm]})
                    if populated:
                        row_shapes["-".join(map(str, populated))] += 1
                last_col = max((openpyxl.utils.column_index_from_string(c) for c in cols), default=1)
                dimension = f"A1:{openpyxl.utils.get_column_letter(last_col)}{max(used_rows, default=1)}"
                book["sheets"].append({"index": index, "dimension": dimension, "nonemptyRows": len(used_rows), "cellTypes": dict(counts), "headerCandidates": headers, "dateHeaderCells": dates, "columnTypes": dict(cols), "distinctRowShapes": len(row_shapes), "visibility": ws.sheet_state})
            wb.close()
            books.append(book)
    return books


if __name__ == "__main__":
    if "--snapshot" in sys.argv:
        write("SOURCE_BASELINE.json", source_snapshot())
        write("BRANCH_REFRESH.json", branches())
        print("Source identity and fetched branch comparison recorded.")
    elif "--branches" in sys.argv:
        write("BRANCH_REFRESH.json", branches())
        print("Branch path inventory refreshed without replacing the source baseline.")
    elif "--worksheets" in sys.argv:
        books = worksheets()
        write("WORKSHEET_EVIDENCE.json", books)
        for b in books:
            print(b["file"])
            for s in b["sheets"]:
                print(s["index"], s["dimension"], "nonemptyRows", s["nonemptyRows"], "types", s["cellTypes"], "headers", [h for h in s["headerCandidates"] if re.fullmatch(r"[A-Z]+[1-3]", h["cell"])], "dateHeaders", len(s["dateHeaderCells"]))
    elif "--verify-preservation" in sys.argv:
        before = json.loads((OUT / "SOURCE_BASELINE.json").read_text(encoding="utf-8"))
        after = source_snapshot()
        delta = sorted(k for k in set(before["files"]) | set(after["files"]) if before["files"].get(k) != after["files"].get(k))
        result = {"passed": not delta and before["head"] == after["head"], "fileCount": len(after["files"]), "changedPaths": delta, "headUnchanged": before["head"] == after["head"]}
        write("PRESERVATION_CHECK.json", result)
        print(json.dumps(result))
        sys.exit(0 if result["passed"] else 1)
